from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image

OUT = Path("attached_assets/entrega_vrf")
OUT.mkdir(parents=True, exist_ok=True)
PUBLIC = Path("artifacts/vrf-project-video/public")
r = Presentation()
r.slide_width, r.slide_height = Inches(9), Inches(16)
WHITE, ORANGE, GRAY = "FFFFFF", "FF6200", "B7BBC4"

def text(s, value, x, y, w, h, size=25, color=WHITE, bold=False):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(14)
    return box

def card(s, x, y, w, h):
    sh = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string("191B20")
    sh.line.color.rgb = RGBColor.from_string("33353A")

def picture(s, path, x, y, w, h):
    with Image.open(path) as im:
        ratio = min(w / im.width, h / im.height)
        iw, ih = im.width * ratio, im.height * ratio
    s.shapes.add_picture(str(path), Inches(x+(w-iw)/2), Inches(y+(h-ih)/2),
                         width=Inches(iw), height=Inches(ih))

def slide(n, title, subtitle, duration):
    s = r.slides.add_slide(r.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string("0C0D10")
    text(s, "HARD SERVICE  /  MONITOR VRF LG", .65, .55, 7.7, .6, 15, ORANGE, True)
    text(s, title, .65, 1.6, 7.7, 2.1, 40, WHITE, True)
    text(s, subtitle, .65, 3.8, 7.7, 2.0, 24, GRAY)
    text(s, f"{n:02d} / 06", .65, 15, 3, .45, 14, GRAY)
    t = OxmlElement("p:transition")
    t.set("advClick", "1")
    t.set("advTm", str(duration*1000))
    t.append(OxmlElement("p:fade"))
    s._element.append(t)
    s.notes_slide.notes_text_frame.text = (
        f"Duração sugerida: {duration} segundos. Textos e formas são editáveis. "
        "Imagens são substituíveis. Adaptação em slides do vídeo original, "
        "não preserva suas animações React. Para gerar vídeo no PowerPoint desktop: "
        "Arquivo > Exportar > Criar um Vídeo. Use os intervalos gravados."
    )
    return s

s=slide(1,"Monitoramento\nVRF LG","Programa desenvolvido pelo time Hard Service para monitorar a saúde dos equipamentos.",6)
picture(s,PUBLIC/"images/xp-logo.jpg",1.5,6.3,6,4)
text(s,"Leituras periódicas.\nAcompanhamento da saúde.\nPrevenção de indisponibilidade.",.8,11,7.4,2.6,27)
s=slide(2,"Manutenção\npreditiva","Análise dos parâmetros extraídos dos equipamentos VRF com o software LGMV.",8)
for y, head, body in [(6.3,"01  LEITURA LGMV","Parâmetros operacionais registrados periodicamente."),
                      (9,"02  COMPARAÇÃO","Relatório de startup e baseline da LG."),
                      (11.7,"03  ACOMPANHAMENTO","Identificação de desvios para apoiar a manutenção.")]:
    card(s,.65,y,7.7,2.2); text(s,head,.95,y+.2,7.1,.55,21,ORANGE,True)
    text(s,body,.95,y+.85,7.1,1.1,22)
s=slide(3,"Parâmetros\ne insights","Exemplo de análise registrada no programa. Leituras não estabilizadas em refrigeração.",8)
for y, label, value in [(6.2,"PRESSÃO DE DESCARGA","221,62 kPa"),(8.2,"PRESSÃO DE SUCÇÃO","213,35 kPa"),(10.2,"EEV","40 pulsos")]:
    card(s,.65,y,7.7,1.7); text(s,label,.95,y+.15,7,.5,17,GRAY)
    text(s,value,.95,y+.7,7,.75,30,ORANGE,True)
text(s,"INSIGHT\nPressões e EEV anômalas: repetir a coleta LGMV após estabilização.",.85,12.5,7.3,2,23)
s=slide(4,"Referência\ntécnica LG","Na ausência do relatório de startup, a tabela LG apoia a avaliação de sistemas com condensação a ar.",8)
picture(s,PUBLIC/"images/lg-reference-table.png",.65,6.1,7.7,5.7)
text(s,"Aplicação: Air-cooled.\nNão substituir o baseline específico de sistemas com condensação a água.",.85,12.4,7.3,2,23,GRAY)
s=slide(5,"Relatório\ntécnico","Diagnóstico, insights e plano de ação — com exportação para Word.",8)
picture(s,PUBLIC/"diagnostico-real.jpg",.65,6,7.7,6.1)
text(s,"DIAGNÓSTICO\nLeituras não estabilizadas.",.85,12.1,7.3,1,21,ORANGE,True)
text(s,"PLANO DE AÇÃO\nRepetir a coleta após estabilização.",.85,13.35,7.3,1.1,21)
s=slide(6,"Prevenção\nativa.","Um processo de Hard Service para acompanhar a saúde dos sistemas VRF LG.",4)
picture(s,PUBLIC/"images/xp-logo.jpg",1.5,6.5,6,4)
text(s,"Dados para orientar ações.\nCuidado contínuo com os equipamentos.",.85,11.4,7.3,2,28)
target=OUT/"VRF_Monitor_Editavel.pptx"
r.save(target)
# Reopen the output and verify editable content and dimensions.
check=Presentation(target)
assert len(check.slides)==6
assert all(any(sh.has_text_frame for sh in sl.shapes) for sl in check.slides)
assert check.slide_width/check.slide_height == 9/16
print(target)